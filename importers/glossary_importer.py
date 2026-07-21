import re
from collections import defaultdict, Counter
from datetime import datetime

from database.database import get_connection
from importers.tm_importer import normalize_segment, _upsert_rows


BRAND_PATTERNS = [
    r"^[A-Z][a-z]+[A-Z]",
    r"\b(IKEA|Home24|Wayfair|Amazon)\b",
]

COLLECTION_SUFFIXES = [r"\b\w+ (Serie|Collection|Kollektion|Line)\b"]

MIN_TERM_LENGTH = 3
MIN_FREQUENCY = 2


def build_glossary_from_tm(progress_callback=None) -> dict:
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT source_segment, target_segment, frequency FROM translation_memory"
        ).fetchall()

    term_pairs: dict[str, Counter] = defaultdict(Counter)

    for i, row in enumerate(rows):
        source = row["source_segment"]
        target = row["target_segment"]
        freq = row["frequency"] or 1

        src_parts = _split_segment(source)
        tgt_parts = _split_segment(target)

        if len(src_parts) == 1 and len(tgt_parts) == 1:
            src_term = src_parts[0].strip()
            tgt_term = tgt_parts[0].strip()
            if _is_valid_term(src_term) and _is_valid_term(tgt_term):
                term_pairs[src_term.lower()][tgt_term] += freq

        if progress_callback and i % 5000 == 0:
            progress_callback(i / len(rows))

    inserted = 0
    skipped = 0

    with get_connection() as conn:
        for src_lower, target_counter in term_pairs.items():
            if sum(target_counter.values()) < MIN_FREQUENCY:
                skipped += 1
                continue

            best_target, best_freq = target_counter.most_common(1)[0]
            category = _detect_term_category(src_lower)
            confidence = min(1.0, best_freq / max(sum(target_counter.values()), 1))

            try:
                conn.execute(
                    "INSERT OR IGNORE INTO glossary "
                    "(source_term, target_term, category, frequency, confidence, source_type) "
                    "VALUES (?,?,?,?,?,'TM')",
                    (src_lower, best_target, category, best_freq, confidence),
                )
                inserted += 1
            except Exception:
                skipped += 1

    return {"inserted": inserted, "skipped": skipped}


def _split_segment(text: str) -> list[str]:
    parts = re.split(r"<br\s*/?>|,\s*(?=[A-ZÜÄÖ])|/(?=[A-ZÜÄÖ])|(?<=\w)\s+-\s+(?=\w)", text)
    return [p.strip() for p in parts if p.strip()]


def _is_valid_term(term: str) -> bool:
    if len(term) < MIN_TERM_LENGTH:
        return False
    if re.match(r"^\d+$", term):
        return False
    for pattern in BRAND_PATTERNS:
        if re.search(pattern, term):
            return False
    for pattern in COLLECTION_SUFFIXES:
        if re.search(pattern, term, re.IGNORECASE):
            return False
    if re.match(r"^[A-Z][a-z]+ [A-Z][a-z]+$", term):
        return False
    return True


TERM_CATEGORIES = {
    "kitchen": ["küche", "keuken", "herd", "kochen"],
    "bathroom": ["bad", "dusch", "wanne", "sanitär"],
    "bedroom": ["bett", "schlaf", "matratze"],
    "color": ["farb", "weiss", "schwarz", "blau", "rot", "grün", "grau", "braun", "beige"],
    "material": ["holz", "metall", "stoff", "leder", "glas", "eiche", "buche", "kiefer", "mdf"],
    "furniture": ["sofa", "tisch", "stuhl", "schrank", "regal", "kommode"],
    "lighting": ["lampe", "leuchte", "licht", "led"],
    "textile": ["kissen", "decke", "vorhang", "teppich"],
}


def _detect_term_category(term: str) -> str:
    t = term.lower()
    for cat, kws in TERM_CATEGORIES.items():
        if any(kw in t for kw in kws):
            return cat
    return "general"


def import_glossary_from_excel(filepath: str) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).lower() if c else "" for c in rows[0]]

    src_col = next((i for i, h in enumerate(header) if "source" in h or "de" in h or "german" in h), 0)
    tgt_col = next((i for i, h in enumerate(header) if "target" in h or "nl" in h or "dutch" in h), 1)
    cat_col = next((i for i, h in enumerate(header) if "cat" in h), None)

    inserted = 0
    with get_connection() as conn:
        for row in rows[1:]:
            src = str(row[src_col]).strip() if row[src_col] else None
            tgt = str(row[tgt_col]).strip() if row[tgt_col] else None
            cat = str(row[cat_col]).strip() if cat_col is not None and row[cat_col] else "general"
            if not src or not tgt:
                continue
            try:
                conn.execute(
                    "INSERT OR IGNORE INTO glossary (source_term, target_term, category, source_type) VALUES (?,?,?,'IMPORTED')",
                    (src.lower(), tgt, cat),
                )
                inserted += 1
            except Exception:
                pass

    return {"inserted": inserted}


# ── Official DE→NL glossary import (authoritative source, PART 1 §9) ───────
#
# Rows are classified into three destinations rather than dumped into one
# table, reusing the machinery each destination already has:
#   - "Label:" rows        -> glossary table (source_type=OFFICIAL_GLOSSARY);
#                              terminology.py's colon-label layer picks these up.
#   - 2-word "Type Model"  -> translation_memory, via the SAME upsert path
#     rows (e.g. "Klapp-      importers/tm_importer.py already uses. AdaptiveTM
#     tisch Raza")            extracts the Klapptisch {MODEL} -> Klaptafel
#                              {MODEL} pattern from this at lookup time — no new
#                              matching logic needed.
#   - everything else       -> glossary table (general vocabulary).
#
# "DE == NL" rows (no information) and case-insensitive duplicate source terms
# (kept: first occurrence only) are skipped and counted, not silently dropped.

def _pick_glossary_sheet(wb):
    for name in wb.sheetnames:
        low = name.lower()
        if "glossary" in low or "technical" in low or "woordenlijst" in low:
            return wb[name]
    return wb.active


def _looks_like_product_model(source: str, target: str) -> bool:
    """True when a 2-word DE term is "ProductType ModelName" rather than an
    ordinary 2-word phrase (e.g. "Akazie Hell" material+color, or "Beleuchteter
    Spiegel" adjective+noun).

    Model-shaped-token alone (ModelNameProtector's heuristic — any capitalized
    word outside its small hardcoded vocabulary) is far too broad here: German
    capitalizes every noun, so most ordinary 2-word phrases in a 14k-row
    glossary would false-positive. The reliable signal is that the DATA itself
    already proves it — a real model name is never translated, so it survives
    as the target's last word too (e.g. "Klapptisch Raza" -> "Klaptafel Raza").
    """
    words = source.split()
    if len(words) != 2:
        return False
    model = words[1]
    if not (model[:1].isupper() or any(ch.isdigit() for ch in model)):
        return False
    tgt_words = target.split()
    return bool(tgt_words) and tgt_words[-1] == model


def import_official_glossary(file_bytes: bytes, progress_callback=None) -> dict:
    """Import the official DE→NL furniture glossary from raw .xlsx bytes.

    Idempotent — safe to re-run when the user supplies an updated file.
    Returns a stats dict: total, de_eq_nl_skipped, duplicates_skipped
    (with the list of skipped (source, first_target, dropped_target) triples
    under 'conflicts'), labels_inserted, general_inserted, tm_pairs_inserted.
    """
    import io
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = _pick_glossary_sheet(wb)
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return {"total": 0}

    # Locate the DE/NL columns by header if present, else assume col 0/1.
    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    src_col = next((i for i, h in enumerate(header) if "german" in h or h in ("de", "duits")), 0)
    tgt_col = next((i for i, h in enumerate(header) if "dutch" in h or h in ("nl", "nederlands")), 1)
    data_rows = rows[1:] if header and (header[src_col] or header[tgt_col]) else rows

    total = len(data_rows)
    seen: dict[str, str] = {}          # lowercased source -> first target kept
    conflicts: list[tuple[str, str, str]] = []
    de_eq_nl = 0
    glossary_rows: list[tuple[str, str, str]] = []   # (source_term_lower, target_term, category)
    tm_pairs: list[dict] = []

    for i, row in enumerate(data_rows):
        src = str(row[src_col]).strip() if src_col < len(row) and row[src_col] else ""
        tgt = str(row[tgt_col]).strip() if tgt_col < len(row) and row[tgt_col] else ""
        if progress_callback and i % 1000 == 0:
            progress_callback(i / total if total else 1.0)
        if not src or not tgt:
            continue
        if normalize_segment(src) == normalize_segment(tgt):
            de_eq_nl += 1
            continue

        key = src.lower()
        if key in seen:
            conflicts.append((src, seen[key], tgt))
            continue
        seen[key] = tgt

        if src.endswith(":"):
            glossary_rows.append((key, tgt, "label"))
        elif _looks_like_product_model(src, tgt):
            tm_pairs.append({
                "source": src, "target": tgt, "freq": 100,
                "created": None, "created_by": "OFFICIAL_GLOSSARY", "source_id": None,
            })
        else:
            glossary_rows.append((key, tgt, _detect_term_category(key)))

    labels_inserted = general_inserted = 0
    with get_connection() as conn:
        for source_term, target_term, category in glossary_rows:
            try:
                cur = conn.execute(
                    "INSERT OR IGNORE INTO glossary "
                    "(source_term, target_term, category, frequency, confidence, source_type, active) "
                    "VALUES (?,?,?,100,0.98,'OFFICIAL_GLOSSARY',1)",
                    (source_term, target_term, category),
                )
                if cur.rowcount:
                    if category == "label":
                        labels_inserted += 1
                    else:
                        general_inserted += 1
            except Exception:
                pass

    tm_stats = _upsert_rows(tm_pairs, progress_callback=None) if tm_pairs else {"inserted": 0, "updated": 0}

    # Refresh in-memory caches so the import applies immediately.
    try:
        from engines.nl.terminology import get_terminology
        get_terminology().reload()
    except Exception:
        pass
    try:
        from engines.nl.model_protector import get_model_protector
        get_model_protector().reload()
    except Exception:
        pass

    if progress_callback:
        progress_callback(1.0)

    return {
        "total": total,
        "de_eq_nl_skipped": de_eq_nl,
        "duplicates_skipped": len(conflicts),
        "conflicts": conflicts,
        "labels_inserted": labels_inserted,
        "general_inserted": general_inserted,
        "tm_pairs_inserted": tm_stats.get("inserted", 0),
        "tm_pairs_updated": tm_stats.get("updated", 0),
    }
